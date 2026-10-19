"""Typeset the four-page construction plan using the bundled python-docx runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
FOLDER=ROOT/'data/demo/s1_case_001/documents'
doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.5); section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(0.7)
section.left_margin=section.right_margin=Inches(0.75)
for name,size in [('Normal',11),('Title',21),('Heading 1',15),('Heading 2',12)]:
    style=doc.styles[name]
    style.font.name='Microsoft YaHei'; style.font.size=Pt(size); style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    style.paragraph_format.line_spacing=1.2
    style.paragraph_format.space_after=Pt(6)
    if name.startswith('Heading'):
        style.paragraph_format.space_before=Pt(12)
        style.paragraph_format.keep_with_next=True
    borders=style.element.find('.//'+qn('w:pBdr'))
    if borders is not None: borders.getparent().remove(borders)
doc.core_properties.title='四层A区洞口防护与作业协调施工方案'
doc.core_properties.author='BIM Safety Coordination Demo'


def table(rows):
    cols=len(rows[0]); t=doc.add_table(rows=1, cols=cols)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    widths=([1.15,1.25,4.6] if cols==3 else [2.35,1.05,1.25,2.35])
    if cols==4 and rows[0][0]=='作业': widths=[2.05,1.55,1.55,1.85]
    if cols==3 and '检查项' in rows[0][0]: widths=[1.9,2.9,2.2]
    for i,w in enumerate(widths): t.columns[i].width=Inches(w)
    for n,row in enumerate(rows):
        cells=t.rows[0].cells if n==0 else t.add_row().cells
        for i,text in enumerate(row):
            cell=cells[i]; cell.width=Inches(widths[i]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=text; tcpr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                x=OxmlElement('w:'+edge); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); borders.append(x)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                x=OxmlElement('w:'+edge); x.set(qn('w:w'),'85'); x.set(qn('w:type'),'dxa'); margins.append(x)
            tcpr.append(margins)
            if n==0:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EFF6'); tcpr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.15
                for run in p.runs:
                    run.font.size=Pt(10.5); run.bold=(n==0)
            no_split=OxmlElement('w:cantSplit'); t.rows[n]._tr.get_or_add_trPr().append(no_split)
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)


lines=(FOLDER/'区域施工方案.md').read_text(encoding='utf-8').splitlines()
pending=[]
for line in lines+['']:
    if line.startswith('|'):
        row=[x.strip() for x in line.strip('|').split('|')]
        if not all(set(x)<=set('- :') for x in row): pending.append(row)
        continue
    if pending: table(pending); pending=[]
    if not line.strip(): continue
    if line=='---PAGE---': doc.add_page_break()
    elif line.startswith('# '): doc.add_paragraph(line[2:],style='Title')
    elif line.startswith('## '): doc.add_paragraph(line[3:],style='Heading 1')
    else: doc.add_paragraph(line)
path=FOLDER/'区域施工方案.docx'
doc.save(path)
print(path)
