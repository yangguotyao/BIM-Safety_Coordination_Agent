from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "00_执行计划与自检/资料整合清单.xlsx"
TEAL = "0F766E"
DARK = "0B1728"
LIGHT = "E6F4F1"
AMBER = "FEF3C7"
RED = "FEE2E2"
GRAY = "F1F5F9"

wb = Workbook()
ws = wb.active
ws.title = "总览"

header_font = Font(name="PingFang SC", bold=True, color="FFFFFF", size=11)
body_font = Font(name="PingFang SC", size=10)
title_font = Font(name="PingFang SC", bold=True, size=18, color="0B1728")
sub_font = Font(name="PingFang SC", size=10, color="475569")
thin = Side(style="thin", color="CBD5E1")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_sheet(sheet, widths, header_row=1, wrap_cols=None, freeze="A2", autofilter=True):
    for idx, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    for cell in sheet[header_row]:
        cell.font = header_font
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = border
    sheet.row_dimensions[header_row].height = 26
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=sheet.max_row, max_col=len(widths)):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=wrap_cols is None or cell.column in wrap_cols)
            cell.border = border
    if freeze:
        sheet.freeze_panes = freeze
    if autofilter and sheet.max_row > header_row:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(widths))}{sheet.max_row}"

# ---------------- 总览 ----------------
ws.merge_cells("A1:F2")
ws["A1"] = "BIM 施工安全协同平台 · 资料整合清单"
ws["A1"].font = title_font
ws["A1"].alignment = Alignment(vertical="center")
ws.merge_cells("A3:F3")
ws["A3"] = "依据：原始方案 DOCX、7 张聊天截图、06_Astra任务书、数据字典、SafetyCase Schema、12 项验收用例 · 2026-10-18"
ws["A3"].font = sub_font

rows = [
    ("交付包", "BIM-Safety-Coordination-Agent", ""),
    ("案例名称", "BIM-Safety Coordination Agent / IM-Safety Coordination Agent", ""),
    ("核心主张", "不是让 LLM 操作 BIM，而是让 Agent 把施工安全任务、BIM 空间信息、施工进度和安全规范连接成可复核 Safety Case。", ""),
    ("主数据", "S1 bSDD_references_example.ifc，394,697,402 字节，SHA256 4075c44a6e4769b10a0f3ef2c925847899c76444ea33b81d3caffa85f321f9f2", ""),
    ("真实部分", "IFC 板/墙几何、GUID、19 楼层、16,372 构件、1,247 个 05 层索引、35 个预览网格、官方规范 PDF", ""),
    ("模拟部分", "人员位置、防护状态、施工活动、2.5m 坠落高度假设、1.0m 暴露策略、12 张演示示意图", ""),
    ("不宣称", "真实事故概率、现场合规、完整模型全量、真实 MEP/洞口/4D 识别", ""),
    ("技术栈", "Vue3 + TypeScript + Vite + Three.js / FastAPI + Pydantic + LangGraph + SQLite + IfcOpenShell", ""),
    ("运行地址", "http://127.0.0.1:8000", ""),
    ("测试状态", "pytest 9 passed；12/12 验收通过；Playwright 全流程通过；0 console/page error", ""),
    ("包用途", "文献综述案例展示、BIM 施工安全 Agent 研究演示、可复核闭环验证", ""),
]
ws.append([])
for row in rows:
    ws.append(row)
for r in range(5, ws.max_row + 1):
    ws.cell(r, 1).font = Font(name="PingFang SC", bold=True, size=10, color="0F766E")
    ws.cell(r, 1).alignment = Alignment(vertical="top", wrap_text=True)
    ws.cell(r, 2).font = body_font
    ws.cell(r, 2).alignment = Alignment(vertical="top", wrap_text=True)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    ws.cell(r, 1).border = border
    ws.cell(r, 2).border = border
ws.column_dimensions["A"].width = 18
for col in "BCDEF":
    ws.column_dimensions[col].width = 26
ws.sheet_view.showGridLines = False

# ---------------- 资料清单 ----------------
ws2 = wb.create_sheet("资料清单")
ws2.append(["类别", "资料 / 文件", "相对路径", "用途", "取得/核验状态", "是否阻断首版"])
items = [
    ("原始方案", "BIM-Safety_Coordination_Agent_方案 (1).docx", "01_资料与依据/02_原始附件/", "案例目标、六类资料、Safety Case、工具层、Agent 角色、13 步流程", "已取得，全文与 3 张表已读取", "否"),
    ("原始附件", "聊天截图_01~07.png", "01_资料与依据/02_原始附件/", "用户原始诉求、参考平台、目录示意、数据集建议", "已取得并核对", "否"),
    ("任务书", "一次性完成Demo任务书.md", "01_资料与依据/06_Astra任务书/", "实现范围、数据限制、验收标准", "已读取", "否"),
    ("接口", "接口与工具约定.md", "01_资料与依据/06_Astra任务书/", "HTTP/错误/工具合同/SQLite", "已读取并落实", "否"),
    ("数据字典", "数据字典与状态约定.md", "01_资料与依据/07_案例与验收模板/", "状态机、版本、时间、provenance", "已读取并落实", "否"),
    ("Schema", "SafetyCase.schema.json", "01_资料与依据/07_案例与验收模板/", "案例结构", "已读取", "否"),
    ("样例", "SafetyCase样例.json / 完整闭环案例参考.json", "01_资料与依据/07_案例与验收模板/", "闭环参考", "已读取", "否"),
    ("样例", "施工任务样例.json / 检查区域.json", "01_资料与依据/07_案例与验收模板/", "任务、区域、边线、人员点", "已读取", "否"),
    ("证据", "演示证据图 EV-001~012.png", "01_资料与依据/07_案例与验收模板/演示证据图/", "12 条模拟观察", "12 张全部存在，1200×800", "否"),
    ("验收", "验收用例.json / 验收用例.csv", "01_资料与依据/07_案例与验收模板/", "12 项 oracle", "已读取，执行器逐项使用", "否"),
    ("规范", "规范条款种子.json", "01_资料与依据/04_安全规范/", "7 active + 2 retired", "已读取，FTS5 索引", "否"),
    ("规范", "GB55034-2022_官方原文.pdf / 发布公告.html", "01_资料与依据/04_安全规范/", "GB 条文释义与原文", "已核验，PDF 5-6 页", "否"),
    ("规范", "JGJ80-2016_官方原文.pdf / 发布公告.html", "01_资料与依据/04_安全规范/", "JGJ 4.3.1 释义与原文", "已核验，PDF 14 页", "否"),
    ("规范", "GB55034-2022_OCR待复核.txt / JGJ80-2016_OCR待复核.txt", "01_资料与依据/04_安全规范/", "工作文本，未全量入 RAG", "保留，按说明不入索引", "否"),
    ("BIM", "bSDD_references_example.ifc", "01_资料与依据/03_数据集/S1/", "主模型、真实几何查询源", "已取得，394,697,402 字节，SHA256 已核验", "否"),
    ("BIM", "区域几何预览.json", "01_资料与依据/03_数据集/S1/解析结果/", "35 个真实网格", "已取得，1,200 顶点 / 2,673 三角面", "否"),
    ("BIM", "选定楼层构件索引.csv", "01_资料与依据/03_数据集/S1/解析结果/", "05 层 1,247 构件", "已取得", "否"),
    ("BIM", "模型检查结果.json", "01_资料与依据/03_数据集/S1/解析结果/", "楼层/构件/限制", "已读取", "否"),
    ("测试模型", "buildingSMART IFC2x3/IFC4/IFC4x3", "01_资料与依据/03_数据集/buildingSMART/", "IFC 解析与洞口测试", "3 个原件已取得", "否"),
    ("备选", "GNI 官方说明/元数据/许可", "01_资料与依据/03_数据集/GNI/", "备选来源说明", "元数据已归档，原件未取得（Zenodo 403）", "否"),
    ("备选", "RestageIFC 官方元数据", "01_资料与依据/03_数据集/RestageIFC/", "独立进度参考", "元数据已归档，原件未取得（Zenodo 403）", "否"),
    ("文献", "论文全文 P06/P08/P09/P10/P11/P12", "01_资料与依据/05_文献与技术资料/论文全文/", "Background / RAG / Agent 方法", "6 份 PDF 可打开", "否"),
    ("文献", "论文索引/题录/摘要", "01_资料与依据/05_文献与技术资料/", "12 篇阅读清单与核验", "P01-P05/P07 摘要/题录级；不冒充全文", "否"),
    ("参考", "原项目 README / 参考平台分析", "01_资料与依据/05_文献与技术资料/参考平台/", "形式参考，未复制代码", "已读取", "否"),
    ("缺失", "高处作业全过程风险管控智能体_LangGraph框架介绍.docx", "未提供", "背景附件", "未提供；不推测正文", "否"),
    ("缺失", "S1 现场照片/图纸/点云/4D", "Trimble 分享页", "真实现场证据", "未取得；以模拟观察替代", "否"),
]
for item in items:
    ws2.append(item)
style_sheet(ws2, [12, 34, 42, 50, 38, 14], wrap_cols=None)
for r in range(2, ws2.max_row + 1):
    status = str(ws2.cell(r, 5).value)
    if "未取得" in status or "未提供" in status:
        ws2.cell(r, 5).fill = PatternFill("solid", fgColor=AMBER)
    if ws2.cell(r, 6).value == "否":
        ws2.cell(r, 6).fill = PatternFill("solid", fgColor=LIGHT)

# ---------------- 数据事实 ----------------
ws3 = wb.create_sheet("数据事实")
ws3.append(["事实", "值", "来源 / 说明"])
facts = [
    ("主模型", "bSDD_references_example.ifc", "S1 社区仓库"),
    ("模型字节数", "394,697,402", "文件实际大小"),
    ("SHA256", "4075c44a6e4769b10a0f3ef2c925847899c76444ea33b81d3caffa85f321f9f2", "模型检查结果"),
    ("Schema", "IFC2X3", "IfcOpenShell 0.9.0.post1"),
    ("IfcElement", "16,372", "模型检查结果"),
    ("IfcBuildingStorey", "19", "模型检查结果"),
    ("IfcSlab / IfcWall", "573 / 7,426", "模型检查结果"),
    ("IfcOpeningElement", "0", "无显式洞口；不等于几何无洞口"),
    ("IfcTask / IfcWorkSchedule", "0 / 0", "不能声称真实 4D"),
    ("案例楼层", "05 vijfde verdieping", "GUID 1Vhy2b2Mj9FRewx9CZZoux"),
    ("05 层索引", "1,247 构件", "选定楼层构件索引.csv"),
    ("首屏网格", "35 对象（4 Slab + 31 Wall）", "区域几何预览.json"),
    ("主构件", "0FOHK4aT910wgskSbVHn0U", "IfcSlab"),
    ("世界 Z", "17.300 m", "不是坠落高度"),
    ("演示坠落高度", "2.5 m", "模拟假设"),
    ("暴露阈值", "1.0 m", "演示项目策略，不是国标"),
    ("规范", "GB 55034-2022 / JGJ 80-2016", "7 active + 2 retired"),
    ("证据", "12 张，1200×800", "真实 IFC 几何 + 模拟标记"),
    ("固定时钟", "2026-10-16T09:00:00+08:00", "演示默认"),
]
for f in facts:
    ws3.append(f)
style_sheet(ws3, [26, 60, 50])

# ---------------- 需求-实现 ----------------
ws4 = wb.create_sheet("需求-实现映射")
ws4.append(["编号", "任务书要求", "实现位置", "验证", "状态"])
reqs = [
    ("D1", "Vue3 + TS + Vite + Three.js", "frontend/src", "构建通过 + 浏览器", "完成"),
    ("D2", "FastAPI + Pydantic + IfcOpenShell + LangGraph + SQLite", "backend/app", "pytest + 运行", "完成"),
    ("D3", "3D 旋转/缩放/点选/GUID/类别/楼层/风险高亮/证据定位", "BimView + BimViewer", "Playwright + 截图", "完成"),
    ("D4", "米制、局部原点、GUID 映射", "BimViewer.toThree / tools.py", "代码 + 截图", "完成"),
    ("D5", "超出范围告知 + 按需加载", "/elements/{guid}/mesh + scope note", "10.75s 首次提取", "完成"),
    ("D6", "案例详情完整字段", "CaseDetailView 5 tabs", "Playwright", "完成"),
    ("D7", "新建/分析/补证/审核/实施/复核/重开/导出", "services.py + API", "pytest + UI 全流程", "完成"),
    ("D8", "持久化、刷新不丢", "SQLite + WAL", "重启后状态一致", "完成"),
    ("D9", "12 验收 + 重置", "acceptance.py + AcceptanceView", "12/12", "完成"),
    ("E1", "六类职责", "agent.py 10 节点", "trace 截图", "完成"),
    ("E2", "确定性几何/时间工具", "tools.py", "pytest", "完成"),
    ("E3", "AABB 仅候选", "spatial_check", "代码/输出", "完成"),
    ("E4", "unknown/needs_evidence", "evidence_check / risk_assessment", "T04/T05", "完成"),
    ("E5", "工具统一字段", "tools.py", "trace", "完成"),
    ("E6", "禁止任意代码", "无 exec/eval", "代码审查", "完成"),
    ("E7", "真实 trace", "agent_events + AgentView", "UI", "完成"),
    ("E8", "最多重试一次 + 不完整", "Trace.call + analysis_incomplete", "live 失败测试", "完成"),
    ("F1", "默认 DEMO_MODE=true", "config.py", "health", "完成"),
    ("F2", "无密钥完成全部验收", "deterministic templates", "12/12", "完成"),
    ("F3", "演示模型输出标记", "finding.narrative.label", "UI", "完成"),
    ("F4", "真实模式 OpenAI 兼容", "live_model.py", "代码 + 失败测试", "完成"),
    ("F5", "密钥仅后端", "Authorization header", "代码审查", "完成"),
    ("F6", "配置不完整 503", "start_analysis", "8001 实测", "完成"),
    ("F7", "失败不静默降级", "_llm_reasoning", "8002 实测", "完成"),
    ("F8", "Pydantic + 白名单", "LiveNarrative / _allowed_ids", "代码", "完成"),
    ("G1", "active only", "rules_fts", "T06", "完成"),
    ("G2", "条文释义 + 页码 + 原文", "KnowledgeView", "UI", "完成"),
    ("G3", "OCR 不入 RAG", "seed.py", "代码", "完成"),
    ("G4", "废止强条排除", "retired", "T06", "完成"),
    ("G5", "FTS5 中文二元组 top3", "bigram_tokens", "pytest", "完成"),
    ("G6", "无规则需补依据", "needs_evidence", "T04/T05", "完成"),
    ("G7", "不输出确定合规", "uncertainties", "UI/导出", "完成"),
    ("G8", "演示判定", "risk_assessment", "T01-T03", "完成"),
    ("G9", "no_issue_in_checked_items", "risk_assessment", "T01", "完成"),
    ("G10", "1.2m 只是一项检查", "文案/检查项", "UI", "完成"),
    ("G11", "风险等级演示策略", "UI/导出", "UI", "完成"),
    ("H1", "状态机", "services.py", "pytest", "完成"),
    ("H2", "未批准不得整改", "implement_case", "T11", "完成"),
    ("H3", "人工实施记录", "implementation", "UI", "完成"),
    ("H4", "after 证据复核", "verification_check", "T07/T08", "完成"),
    ("H5", "不能凭建议关闭", "状态机", "T11", "完成"),
    ("H6", "复核失败 reopened", "verify", "T07", "完成"),
    ("H7", "关闭后重开保留历史", "reopen_case", "T12", "完成"),
    ("H8", "修改措施版本", "review modify", "代码", "完成"),
    ("H9", "旧审批不覆盖新方案", "expected_case_version", "代码", "完成"),
    ("H10", "409/422/404", "APIError", "pytest", "完成"),
    ("I1", "start.sh / start.ps1", "scripts", "start.sh 实测", "完成"),
    ("I2", "解析根目录/venv/种子/端口/URL", "start.sh", "启动日志", "完成"),
    ("I3", "离线演示", "本地资源", "浏览器", "完成"),
    ("I4", "单例/缓存", "mesh.py", "10.75s 首次 + 缓存", "完成"),
    ("I5", "固定时钟", "config.py", "UI/API", "完成"),
    ("I6", "真实模式可选检查时刻", "新建案例 as_of", "UI", "完成"),
    ("I7", "全部测试", "pytest/Playwright", "报告", "完成"),
    ("J1", "文档/计划/自检/报告/手册", "交付包", "文件存在", "完成"),
    ("J2", "边界可检索", "README/关于/案例/导出", "文本检索", "完成"),
]
for r in reqs:
    ws4.append(r)
style_sheet(ws4, [8, 44, 40, 32, 10])
for r in range(2, ws4.max_row + 1):
    ws4.cell(r, 5).fill = PatternFill("solid", fgColor=LIGHT)

# ---------------- 缺口与替代 ----------------
ws5 = wb.create_sheet("缺口与替代")
ws5.append(["缺口", "影响", "当前处理", "后续可替换", "是否阻断"])
gaps = [
    ("S1 真实现场照片/视频", "无真实视觉证据", "12 张演示示意图（真实几何 + 模拟标记）", "取得现场照片后替换 evidence 文件与 observation", "否"),
    ("S1 图纸/点云/4D", "无真实 4D 与点云", "模拟施工活动；明确标注", "导入真实 4D/点云并扩展 schedule_query", "否"),
    ("真实 MEP 专业模型", "不能识别真实 MEP 任务", "使用板边检查场景", "导入 MEP 模型与任务", "否"),
    ("RestageIFC 原件", "无独立进度字段验证", "仅保留元数据/入口", "Zenodo 可访问后下载", "否"),
    ("GNI 原件", "无 224 模型备选库", "仅保留元数据/入口", "Zenodo 可访问后下载", "否"),
    ("真实临边/防护高度/坠落基准面", "不能宣称现场合规", "人工标注边线 + 模拟观察 + 不确定性", "现场复核后替换", "否"),
    ("LangGraph 介绍 Word", "缺失背景附件", "不推测正文；以任务书/方案为准", "用户补充后可更新关于页", "否"),
    ("P01-P05/P07 全文", "文献综述深度受限", "保留题录/摘要级并标注", "获取全文后补充", "否"),
    ("真实模型 API 密钥", "无法执行真实供应商调用", "默认演示；已实现真实模式与失败路径", "配置密钥后实测", "否"),
]
for g in gaps:
    ws5.append(g)
style_sheet(ws5, [34, 34, 48, 44, 10], wrap_cols=None)

# ---------------- 交付物 ----------------
ws6 = wb.create_sheet("交付物")
ws6.append(["交付物", "路径", "说明", "状态"])
files = [
    ("网站源码", "02_网站源码/", "FastAPI + Vue3 + Three.js", "可运行"),
    ("构建产物", "02_网站源码/frontend/dist/", "FastAPI 托管", "已构建"),
    ("启动脚本", "02_网站源码/scripts/start.sh / start.ps1", "macOS/Linux + Windows", "已提供"),
    ("Finder 启动", "启动网站.command / 停止网站.command", "双击启动/停止", "已提供"),
    ("执行计划", "00_执行计划与自检/执行计划.docx", "完整计划", "已生成"),
    ("计划自检", "00_执行计划与自检/计划自检记录.md", "计划检查与偏差", "已生成"),
    ("合规矩阵", "00_执行计划与自检/合规矩阵.md", "任务书逐条映射", "已生成"),
    ("资料清单", "00_执行计划与自检/资料整合清单.xlsx", "本文件", "已生成"),
    ("测试报告", "03_测试与验收/验收执行报告.md", "pytest/Playwright/12 项", "已生成"),
    ("验收结果", "03_测试与验收/acceptance-results.json", "12 项实际结果", "已生成"),
    ("演示手册", "04_演示与截图/演示操作手册.docx", "5 分钟演示路线", "已生成"),
    ("页面截图", "04_演示与截图/", "总览/BIM/案例/Agent/验收", "已生成"),
]
for f in files:
    ws6.append(f)
style_sheet(ws6, [22, 58, 52, 14])

# ---------------- 测试结果 ----------------
ws7 = wb.create_sheet("测试结果")
ws7.append(["测试", "命令/方式", "结果", "证据"])
tests = [
    ("后端 pytest", ".venv/bin/python -m pytest tests/test_api.py -q", "9 passed", "03_测试与验收/api-tests.log"),
    ("前端构建", "npm run build", "vue-tsc + vite 通过", "03_测试与验收/frontend-build.log"),
    ("12 项验收", "POST /api/demo/acceptance/run-all", "12/12 passed", "03_测试与验收/acceptance-results.json"),
    ("浏览器全流程", "python ui_playwright.py", "14 checks, 0 errors", "03_测试与验收/ui-test.log"),
    ("模型点选", "Playwright canvas click + inspector", "通过", "03_测试与验收/screenshots/ui-bim.png"),
    ("按需网格", "GET /elements/{guid}/mesh", "10.75s 首次，缓存可用", "03_测试与验收/on-demand-mesh.log"),
    ("重启持久化", "stop → start → GET case", "状态/版本/事件一致", "03_测试与验收/restart-persistence.log"),
    ("无密钥启动", "DEMO_MODE=true", "health mode=demo", "03_测试与验收/验收执行报告.md"),
    ("无效真实配置", "DEMO_MODE=false 缺三项", "503 live_model_config_incomplete", "03_测试与验收/live-config.log"),
    ("真实调用失败", "fake URL + 2s timeout", "run failed + needs_evidence，无静默降级", "03_测试与验收/live-config.log"),
]
for t in tests:
    ws7.append(t)
style_sheet(ws7, [20, 54, 52, 50])

# unify tab colors
for sheet, color in [(ws, "0F766E"), (ws2, "0F766E"), (ws3, "0F766E"), (ws4, "0F766E"), (ws5, "B45309"), (ws6, "0F766E"), (ws7, "0F766E")]:
    sheet.sheet_properties.tabColor = color

wb.save(OUT)
print("saved", OUT)
