from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TEAL = RGBColor(0x0F, 0x76, 0x6E)
DARK = RGBColor(0x0B, 0x17, 0x28)
MUTED = RGBColor(0x47, 0x55, 0x69)
FONT = "PingFang SC"
MONO = "Menlo"


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_margins(cell, top=60, bottom=60, left=90, right=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name=FONT, size=10.5, bold=None, color=None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, size=9, color=MUTED)


def add_bottom_border(paragraph, color="0F766E", size=18) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def clean_inline(text: str) -> str:
    text = text.replace("**", "").replace("`", "")
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    return text.strip()


def add_runs(paragraph, text: str, size=10.5, color=None) -> None:
    # Preserve **bold** segments.
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        value = part[2:-2] if bold else part
        value = value.replace("`", "")
        value = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", value)
        run = paragraph.add_run(value)
        set_run_font(run, size=size, bold=True if bold else None, color=color)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.line_spacing = 1.55
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color, before, after in [
        ("Heading 1", 16, TEAL, 16, 7),
        ("Heading 2", 12.5, DARK, 12, 5),
        ("Heading 3", 11, DARK, 9, 4),
    ]:
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("BIM-Safety Coordination Agent · 交付文档 v1.0 · ")
    set_run_font(run, size=9, color=MUTED)
    add_page_number(footer)
    return doc


def add_cover(doc: Document, title: str, subtitle: str) -> None:
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run("CASE STUDY · 2026")
    set_run_font(run, name=MONO, size=9, color=TEAL, bold=True)
    add_bottom_border(p, size=12)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    run = p.add_run(title)
    set_run_font(run, size=27, color=DARK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run(subtitle)
    set_run_font(run, size=12, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(30)
    run = p.add_run(
        "BIM 施工安全协同平台（BIM-Safety Coordination Agent）\n"
        "围绕 S1 真实 IFC、模拟现场证据、现行规范与 12 项验收，构建可审核、可实施、可复核、可重新打开的 Safety Case。"
    )
    set_run_font(run, size=10.5, color=MUTED)

    doc.add_paragraph()
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Cm(4.2)
    table.columns[1].width = Cm(12.0)
    meta = [
        ("文档版本", "v1.0"),
        ("编制日期", "2026-10-18"),
        ("依据", "原始方案 DOCX、7 张聊天截图、06_Astra任务书、数据字典、SafetyCase Schema、12 项验收"),
        ("交付形态", "访达资料包 + 可运行 Web 平台 + 测试报告 + 演示手册"),
        ("数据边界", "真实 IFC 几何 + 模拟现场观察；不构成现场合规或事故概率结论"),
    ]
    for index, (key, value) in enumerate(meta):
        left, right = table.rows[index].cells
        set_cell_shading(left, "E6F4F1")
        left.text = ""
        right.text = ""
        run = left.paragraphs[0].add_run(key)
        set_run_font(run, size=9.5, bold=True, color=TEAL)
        run = right.paragraphs[0].add_run(value)
        set_run_font(run, size=9.5, color=DARK)
        for cell in (left, right):
            set_cell_margins(cell, 70, 70, 100, 100)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    run = p.add_run("本计划按“完成标志”推进，不设日历工期。执行中如与任务书冲突，以任务书和数据事实为准。")
    set_run_font(run, size=9, color=MUTED)

    doc.add_page_break()


def add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr = table.rows[0].cells
    for index, text in enumerate(header):
        hdr[index].text = ""
        p = hdr[index].paragraphs[0]
        run = p.add_run(clean_inline(text))
        set_run_font(run, size=8.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        set_cell_shading(hdr[index], "0F766E")
        set_cell_margins(hdr[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index in range(len(header)):
            value = row[index] if index < len(row) else ""
            cells[index].text = ""
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            add_runs(p, value, size=8.2, color=DARK)
            if row_index % 2 == 1:
                set_cell_shading(cells[index], "F7FAFC")
            set_cell_margins(cells[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.text = ""
    set_cell_shading(cell, "F1F5F9")
    set_cell_margins(cell, 80, 80, 120, 120)
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.1
    for index, line in enumerate(text.splitlines()):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, name=MONO, size=8, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def convert(md_path: Path, out_path: Path, title: str, subtitle: str) -> None:
    doc = setup_document()
    add_cover(doc, title, subtitle)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_lines: list[str] = []
    table_lines: list[str] = []

    def flush_table() -> None:
        nonlocal table_lines
        if not table_lines:
            return
        rows = []
        header = None
        for line in table_lines:
            if re.match(r"^\s*\|?\s*:?-{3,}", line.replace("|", " ").strip()):
                continue
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if header is None:
                header = cells
            else:
                rows.append(cells)
        if header:
            add_table(doc, header, rows)
        table_lines = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush_table()
            if in_code:
                add_code(doc, "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if stripped.startswith("|"):
            table_lines.append(stripped)
            i += 1
            continue
        flush_table()
        if not stripped:
            i += 1
            continue
        if stripped == "---":
            p = doc.add_paragraph()
            add_bottom_border(p, color="E2E8F0", size=6)
            i += 1
            continue
        if stripped.startswith("# "):
            # Skip the document title; the cover already has it.
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(clean_inline(stripped[3:]), level=1)
        elif stripped.startswith("### "):
            doc.add_heading(clean_inline(stripped[4:]), level=2)
        elif stripped == ">":
            i += 1
            continue
        elif stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            add_runs(p, stripped[1:].strip(), size=10, color=MUTED)
        elif re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            add_runs(p, stripped)
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, stripped[2:])
        elif stripped.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, stripped[2:])
        else:
            p = doc.add_paragraph()
            add_runs(p, stripped)
        i += 1
    flush_table()
    if code_lines:
        add_code(doc, "\n".join(code_lines))

    doc.core_properties.title = title
    doc.core_properties.subject = subtitle
    doc.core_properties.author = "BIM-Safety Coordination Agent"
    doc.core_properties.comments = "Generated from 执行计划.md"
    doc.save(out_path)
    print("saved", out_path)


if __name__ == "__main__":
    convert(
        Path("00_执行计划与自检/执行计划.md"),
        Path("00_执行计划与自检/执行计划.docx"),
        "BIM 施工安全协同 Agent 网站",
        "执行计划 · 资料整合 · 验收与交付",
    )
